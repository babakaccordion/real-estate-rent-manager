from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path
from ..config import DATA_DIR


def build_excel(title: str, sheets: list, filename: str) -> str:
    """sheets: list of (sheet_name, headers, rows). Returns file path."""
    out_dir = DATA_DIR / "exports"
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / filename
    wb = Workbook()
    # remove default
    wb.remove(wb.active)
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="1e3a5f")
    for name, headers, rows in sheets:
        ws = wb.create_sheet(title=name[:31])
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = header_fill
            c.alignment = Alignment(horizontal="right")
            c.border = border
        for row in rows:
            ws.append(row)
        for r in ws.iter_rows(min_row=2):
            for c in r:
                c.alignment = Alignment(horizontal="right")
                c.border = border
        # column widths
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, len(str(h)) + 4)
        ws.freeze_panes = "A2"
    # title sheet? add as first
    wb.save(path)
    return str(path)